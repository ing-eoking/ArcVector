#!/usr/bin/env python3
"""Build the internal-design deck.

    docs/tools/.venv/bin/python docs/tools/deck.py docs/ArcVector-내부설계.pptx
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

W, H = Inches(13.333), Inches(7.5)

INK = RGBColor(0x16, 0x19, 0x1C)
MUTED = RGBColor(0x5B, 0x64, 0x70)
LINE = RGBColor(0xCC, 0xD2, 0xCD)
SOFT = RGBColor(0xE4, 0xE8, 0xE3)
GROUND = RGBColor(0xF1, 0xF3, 0xEF)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
SUNK = RGBColor(0xEA, 0xEC, 0xE7)
ARCUS = RGBColor(0x0D, 0x6B, 0x6B)
USEARCH = RGBColor(0x4A, 0x52, 0xA6)
FAIL = RGBColor(0xA0, 0x36, 0x36)
MOSS = RGBColor(0x4E, 0x6B, 0x34)

SANS = "Apple SD Gothic Neo"
MONO = "Menlo"

MARGIN = Inches(0.85)
BODY_TOP = Inches(2.05)
BODY_W = W - MARGIN * 2


def deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def face(font, name):
    """Set the typeface for latin, East Asian and complex-script runs alike, so
    Korean does not fall back to whatever the viewer picks."""
    font.name = name
    rpr = font._element
    for slot in ("a:ea", "a:cs"):
        for old in rpr.findall(qn(slot)):
            rpr.remove(old)
        el = rpr.makeelement(qn(slot), {"typeface": name})
        rpr.append(el)


def _fill(shape, color):
    if color is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color


def _line(shape, color, width=Pt(1)):
    if color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color
        shape.line.width = width


def text(slide, x, y, w, h, blocks, size=16, color=INK, font=SANS, bold=False,
         align=PP_ALIGN.LEFT, spacing=1.25, anchor=MSO_ANCHOR.TOP, gap=None):
    """`blocks` is a string, or a list of paragraphs; a paragraph is a string or
    a list of (text, {overrides}) runs."""
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0

    if isinstance(blocks, str):
        blocks = [blocks]
    for i, para in enumerate(blocks):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if i:
            p.space_before = Pt(size * 0.55 if gap is None else gap)
        runs = [para] if isinstance(para, str) else para
        for run in runs:
            body, over = (run, {}) if isinstance(run, str) else run
            r = p.add_run()
            r.text = body
            f = r.font
            face(f, over.get("font", font))
            f.size = Pt(over.get("size", size))
            f.bold = over.get("bold", bold)
            f.color.rgb = over.get("color", color)
    return box


def rect(slide, x, y, w, h, fill=SURFACE, edge=LINE, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h
    )
    _fill(shape, fill)
    _line(shape, edge)
    shape.shadow.inherit = False
    if radius:
        shape.adjustments[0] = 0.08
    shape.text_frame.text = ""
    return shape


def label(slide, x, y, w, h, lines, size=13, color=INK, font=SANS, bold=False,
          align=PP_ALIGN.CENTER, spacing=1.15):
    return text(slide, x, y, w, h, lines, size=size, color=color, font=font,
                bold=bold, align=align, spacing=spacing, anchor=MSO_ANCHOR.MIDDLE)


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=Pt(1.25)):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = width
    # python-pptx has no arrowhead API; the head is a child of the line properties.
    ln = c.line._get_or_add_ln()
    head = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(head)
    return c


def rule(slide, y, color=LINE, x=MARGIN, w=None):
    w = BODY_W if w is None else w
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + w, y)
    c.line.color.rgb = color
    c.line.width = Pt(1)
    return c


def code(slide, x, y, w, lines, size=12.5, pad=Inches(0.16), fill=SUNK):
    # Korean falls back to a face with taller metrics than Menlo, so the box is
    # sized above the nominal line height rather than at it.
    h = pad * 2 + Emu(int(Pt(size * 1.62).emu * len(lines)))
    rect(slide, x, y, w, h, fill=fill, edge=LINE)
    text(slide, x + pad, y + pad, w - pad * 2, h - pad * 2, lines,
         size=size, font=MONO, color=INK, spacing=1.25, gap=0)
    return y + h


N = 0


def page(prs, title, eyebrow=None, kicker=None):
    global N
    N += 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = GROUND

    top = Inches(0.62)
    if eyebrow:
        text(slide, MARGIN, top, BODY_W, Inches(0.3), eyebrow.upper(),
             size=11, font=MONO, color=MOSS, bold=True)
        top += Inches(0.36)
    text(slide, MARGIN, top, BODY_W, Inches(0.7), title, size=31, bold=True, color=INK)
    if kicker:
        text(slide, MARGIN, top + Inches(0.62), BODY_W, Inches(0.4), kicker,
             size=15, color=MUTED)
    rule(slide, Inches(1.96))
    text(slide, W - MARGIN - Inches(0.7), H - Inches(0.55), Inches(0.7), Inches(0.3),
         str(N), size=11, font=MONO, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = INK

    text(slide, MARGIN, Inches(2.4), BODY_W, Inches(0.4),
         "ARCVECTOR · 내부 설계", size=13, font=MONO, color=RGBColor(0xA8, 0xC7, 0x85), bold=True)
    text(slide, MARGIN, Inches(2.95), BODY_W, Inches(1.5),
         "두 저장소를 하나처럼 보이게 하는 법",
         size=46, bold=True, color=RGBColor(0xF1, 0xF3, 0xEF))
    text(slide, MARGIN, Inches(4.3), Inches(8.4), Inches(1.4),
         "arcus Map 아이템이 원본이고 usearch 그래프는 캐시다. "
         "명령 하나가 두 저장소를 건드리는데, 밖에서는 언제나 한쪽만 본 것처럼 보여야 한다. "
         "이 발표는 그 제약이 코드를 어떤 모양으로 만들었는지에 대한 것이다.",
         size=16, color=RGBColor(0x94, 0x9D, 0xA8), spacing=1.5)
    return slide


def main(out):
    prs = deck()
    cover(prs)

    # ---------------------------------------------------------------- 1
    s = page(prs, "인덱스 하나가 무엇인가", eyebrow="한 장 요약")
    y = BODY_TOP
    bw, gap = Inches(5.3), Inches(1.13)
    rect(s, MARGIN, y, bw, Inches(2.5), fill=SURFACE, edge=ARCUS)
    label(s, MARGIN + Inches(0.3), y + Inches(0.22), bw - Inches(0.6), Inches(0.34),
          "ARCUS · 원본", size=12, font=MONO, bold=True, color=ARCUS, align=PP_ALIGN.LEFT)
    text(s, MARGIN + Inches(0.3), y + Inches(0.72), bw - Inches(0.6), Inches(1.6),
         ["인덱스 = Map 아이템 하나",
          "벡터 = Map 원소 하나",
          "복제·영속화가 따라가는 것이 이쪽",
          "쓰기가 노드를 떠나는 지점도 이쪽"],
         size=15, spacing=1.5)

    x2 = MARGIN + bw + gap
    rect(s, x2, y, bw, Inches(2.5), fill=SURFACE, edge=USEARCH)
    label(s, x2 + Inches(0.3), y + Inches(0.22), bw - Inches(0.6), Inches(0.34),
          "USEARCH · 캐시", size=12, font=MONO, bold=True, color=USEARCH, align=PP_ALIGN.LEFT)
    text(s, x2 + Inches(0.3), y + Inches(0.72), bw - Inches(0.6), Inches(1.6),
         ["HNSW 노드와 u64 키뿐",
          "복제도 영속화도 안 됨",
          "언제든 Map에서 다시 만들 수 있어야 함",
          "= 잃어도 되는 것"],
         size=15, spacing=1.5)

    arrow(s, MARGIN + bw + Inches(0.12), y + Inches(1.25), x2 - Inches(0.12), y + Inches(1.25),
          color=MUTED, width=Pt(2))
    label(s, MARGIN + bw, y + Inches(0.78), gap, Inches(0.4), "재구축", size=11,
          font=MONO, color=MUTED)

    text(s, MARGIN, y + Inches(2.95), BODY_W, Inches(1.4),
         [[("그래서 규칙이 하나 나온다 — ", {}),
           ("실패는 캐시가 낡는 쪽으로 나야지 원본이 틀리는 쪽으로 나면 안 된다.", {"bold": True})],
          "노드가 하나 남는 것은 무해하다. 아무도 이름을 못 만들어서 모든 조회가 버린다. "
          "반대로 Map에만 남으면 검색에서 영영 안 나오고, 재구축이 없는 빌드는 그것을 복구하지 못한다."],
         size=16, spacing=1.45)

    # ---------------------------------------------------------------- 2
    s = page(prs, "다섯 가지 불변식", eyebrow="이후 모든 결정이 여기서 나온다")
    items = [
        ("I1", "Map이 원본, 그래프는 캐시", "모든 쓰기가 엔진 먼저. 그 반대는 없다"),
        ("I2", "usearch 키 = 그 원소의 주소", "key → id 가 조회가 아니라 역참조. refcount가 그것을 지킨다"),
        ("I3", "원소는 링크된 뒤로 불변", "FILTER가 엔진 락 없이 읽는다. 값이 바뀌면 언제나 새 원소로 교체"),
        ("I4", "실패할 수 있는 일은 CLOG 앞에", "do_map_elem_link가 복제본에 알린다. 그 뒤 실패 = master가 포기한 쓰기를 복제본이 가짐"),
        ("I5", "데몬은 죽지 않는다", "메모리가 없으면 응답으로 답한다. 자라는 컬렉션은 전부 try_reserve"),
    ]
    y = BODY_TOP - Inches(0.1)
    for tag, head, why in items:
        rect(s, MARGIN, y, Inches(0.62), Inches(0.62), fill=INK, edge=None)
        label(s, MARGIN, y, Inches(0.62), Inches(0.62), tag, size=15, font=MONO,
              bold=True, color=GROUND)
        text(s, MARGIN + Inches(0.92), y + Inches(0.02), BODY_W - Inches(0.92), Inches(0.32),
             head, size=18, bold=True)
        text(s, MARGIN + Inches(0.92), y + Inches(0.36), BODY_W - Inches(0.92), Inches(0.3),
             why, size=14, color=MUTED)
        y += Inches(0.94)

    # ---------------------------------------------------------------- 3
    s = page(prs, "원소 하나에 무엇이 들어가나", eyebrow="저장 형식",
             kicker="ATTR은 내용이 얼마든 고정 128바이트. 자리를 아끼지 않는 선택이다.")
    y = BODY_TOP + Inches(0.35)
    segs = [("alen", Inches(1.0), SOFT), ("ATTR  128 B", Inches(3.2), SOFT),
            ("벡터  dim × quant", Inches(6.0), SURFACE), ("\\r\\n", Inches(0.9), SOFT)]
    x = MARGIN
    for name, w, fill in segs:
        rect(s, x, y, w, Inches(0.78), fill=fill, edge=LINE)
        label(s, x, y, w, Inches(0.78), name, size=13, font=MONO)
        x += w
    text(s, MARGIN, y + Inches(0.92), Inches(4.2), Inches(0.3),
         "VECTOR_OFFSET = 130, 상수", size=12, font=MONO, color=MUTED)

    y += Inches(1.55)
    text(s, MARGIN, y, BODY_W, Inches(0.34),
         "벡터는 recovery 빌드에만 들어간다", size=19, bold=True)
    text(s, MARGIN, y + Inches(0.44), Inches(6.2), Inches(1.6),
         ["Map의 벡터를 읽는 곳은 재구축 하나뿐이고, 그건 replication·persistence 빌드에만 있다. "
          "vsim KEY의 질의 벡터는 그래프에서 읽는다.",
          [("768차원 f32 기준 3204 B → 132 B. 2M이면 6GB가 사라진다.", {"bold": True})]],
         size=15, spacing=1.45)

    code(s, MARGIN + Inches(6.7), y + Inches(0.4), BODY_W - Inches(6.7),
         ["#[cfg(recovery)]",
          "pub const fn element_len(&self) -> usize {",
          "    Self::VECTOR_OFFSET + self.vector_bytes()",
          "}",
          "",
          "#[cfg(not(recovery))]",
          "pub const fn element_len(&self) -> usize {",
          "    Self::VECTOR_OFFSET",
          "}"])

    # ---------------------------------------------------------------- 4
    s = page(prs, "키는 원소의 주소다", eyebrow="I2",
             kicker="key → id 가 해시 조회가 아니라 역참조다. 원소가 자기 필드 바이트를 들고 있고, 그게 id다.")
    y = BODY_TOP + Inches(0.3)
    text(s, MARGIN, y, Inches(5.9), Inches(2.4),
         ["한때 여기 id ↔ key 양방향 매핑이 있었다. 슬롯 배열, by_id 해시맵, "
          "인터닝된 문자열, 키에 박은 세대, 슬롯 발급·퇴역·회수.",
          [("2M 항목에 165MB → 주소 집합 40MB 남짓.", {"bold": True})]],
         size=16, spacing=1.45)

    rect(s, MARGIN, y + Inches(2.5), Inches(5.9), Inches(1.5), fill=SURFACE, edge=FAIL)
    text(s, MARGIN + Inches(0.26), y + Inches(2.7), Inches(5.4), Inches(1.1),
         [[("대가는 refcount다. ", {"bold": True, "color": FAIL}),
           ("주소가 주소로 남으려면 원소가 해제되면 안 되고, 그러려면 그래프가 "
            "refcount를 든다. unlink된 원소가 바로 회수되지 않는다.", {})]],
         size=14, spacing=1.4)

    x = MARGIN + Inches(6.5)
    rows = [("주소가 계속 그 원소를 가리킬 것", "refcount"),
            ("역참조 도중 해제되지 않을 것", "지연 회수"),
            ("링크 전 원소가 이름을 갖지 않을 것", "STAGED_TAG"),
            ("밀려난 주소를 놓치지 않을 것", "hold 안에서 읽기"),
            ("인덱스가 사라질 때 전부 놓을 것", "Drop")]
    yy = y
    for what, how in rows:
        text(s, x, yy, Inches(3.6), Inches(0.3), what, size=14)
        text(s, x + Inches(3.7), yy, Inches(1.9), Inches(0.3), how, size=13,
             font=MONO, color=USEARCH, bold=True)
        rule(s, yy + Inches(0.42), color=SOFT, x=x, w=Inches(5.5))
        yy += Inches(0.66)

    # ---------------------------------------------------------------- 5
    s = page(prs, "vadd — 두 저장소에 쓰는 유일한 명령", eyebrow="쓰기",
             kicker="원칙은 하나. 양쪽을 먼저 준비하고, 보이게 만드는 것만 한 락 안에서 한다.")
    steps = [("①", "coords", "본문 → f32", None),
             ("②", "for_write", "메타 읽고 소유권", ARCUS),
             ("③", "reserve_elem", "원소 몸통", ARCUS),
             ("④", "stage", "HNSW add · ~370us", USEARCH),
             ("⑤", "insert", "결과 확정", ARCUS),
             ("⑥", "rename", "주소로 이름", USEARCH),
             ("⑦", "drop_node", "밀어낸 노드", USEARCH)]
    y = BODY_TOP + Inches(0.2)
    w = Inches(1.72)
    x = MARGIN
    for i, (num, name, note, sys) in enumerate(steps):
        held = i in (4, 5)
        rect(s, x, y, w - Inches(0.08), Inches(1.5),
             fill=SUNK if held else SURFACE, edge=sys or LINE)
        text(s, x + Inches(0.14), y + Inches(0.14), w - Inches(0.36), Inches(0.26),
             num, size=13, font=MONO, color=MUTED)
        text(s, x + Inches(0.14), y + Inches(0.46), w - Inches(0.36), Inches(0.4),
             name, size=13.5, font=MONO, bold=True, color=sys or INK)
        text(s, x + Inches(0.14), y + Inches(0.92), w - Inches(0.36), Inches(0.45),
             note, size=11.5, color=MUTED)
        x += w
    rect(s, MARGIN + w * 4, y - Inches(0.34), w * 2 - Inches(0.08), Inches(0.3),
         fill=INK, edge=None)
    label(s, MARGIN + w * 4, y - Inches(0.34), w * 2 - Inches(0.08), Inches(0.3),
          "held 쓰기 락", size=11, font=MONO, bold=True, color=GROUND)

    y += Inches(1.95)
    text(s, MARGIN, y, Inches(6.0), Inches(1.6),
         [[("④가 가장 비싸고, ⑤가 되돌릴 수 없는 지점이다.", {"bold": True})],
          "do_map_elem_link이 CLOG_MAP_ELEM_INSERT를 내보낸다 — 복제본과 영속화 로그가 "
          "이때 이 쓰기를 갖는다. 그래서 실패할 수 있는 일이 전부 그 앞에 있다."],
         size=15, spacing=1.45)
    code(s, MARGIN + Inches(6.6), y - Inches(0.1), BODY_W - Inches(6.6),
         ["// coll_map.c — do_map_elem_link",
          "do_map_elem_replace(info, &pinfo, elem);",
          "CLOG_MAP_ELEM_INSERT(info, old_elem, elem);",
          "//  ↑ 쓰기가 이 노드를 떠나는 지점"])

    # ---------------------------------------------------------------- 6
    s = page(prs, "vcreate는 왜 probe를 하지 않는가", eyebrow="순서",
             kicker="probe는 이미 지나간 순간에 대한 답이다. insert 하나가 네 갈래를 다 알려준다.")
    y = BODY_TOP + Inches(0.15)
    code(s, MARGIN, y, Inches(6.2),
         ["let settled = store",
          "    .alloc_elem(name, META_FIELD, &meta.encode(layout))",
          "    .and_then(|pending| pending.insert_creating(attr));"])
    rows = [("Ok(true)", "이름이 비어 있었다", "Map도 원소도 우리 것 → CREATED"),
            ("Ok(false)", "메타데이터 없는 Map", "우리 것 아님 → 되빼고 안내"),
            ("Err(ElemExists)", "이미 인덱스", "인수하거나 EXISTS"),
            ("Err(BadType)", "Map이 아닌 아이템", "CLIENT_ERROR")]
    yy = y + Inches(1.35)
    for a, b, c in rows:
        text(s, MARGIN, yy, Inches(2.1), Inches(0.3), a, size=13, font=MONO, color=ARCUS, bold=True)
        text(s, MARGIN + Inches(2.2), yy, Inches(2.6), Inches(0.3), b, size=14)
        text(s, MARGIN + Inches(4.9), yy, Inches(6.6), Inches(0.3), c, size=14, color=MUTED)
        rule(s, yy + Inches(0.42), color=SOFT)
        yy += Inches(0.66)

    text(s, MARGIN, yy + Inches(0.18), BODY_W, Inches(0.95),
         [[("엔진의 캐시 락 안에서 정해진 것이라, 보는 것과 하는 것 사이에 창이 없다.", {"bold": True})],
          "등록도 엔진 쓰기보다 앞이다. 그 사이 구간은 '미발행' 표식이 덮는다."],
         size=15, spacing=1.4)

    # ---------------------------------------------------------------- 7
    s = page(prs, "락은 둘, 모양이 다르다", eyebrow="동시성")
    y = BODY_TOP
    for i, (name, what, how, color) in enumerate([
        ("RwLock<Index>", "reserve · reset 배제",
         "공유. usearch가 문서로 요구한다 — 'During reserve, no insertions may be happening'", USEARCH),
        ("RwLock<HeldSet>", "두 저장소 걸친 구간 배제",
         "배타. 구간이 usearch 밖(엔진 호출)까지 간다. 쓰기 경로만", ARCUS),
    ]):
        yy = y + i * Inches(1.35)
        rect(s, MARGIN, yy, BODY_W, Inches(1.15), fill=SURFACE, edge=color)
        text(s, MARGIN + Inches(0.3), yy + Inches(0.18), Inches(3.3), Inches(0.32),
             name, size=15, font=MONO, bold=True, color=color)
        text(s, MARGIN + Inches(3.8), yy + Inches(0.18), Inches(3.2), Inches(0.32),
             what, size=15, bold=True)
        text(s, MARGIN + Inches(0.3), yy + Inches(0.62), BODY_W - Inches(0.6), Inches(0.4),
             how, size=13.5, color=MUTED)

    y += Inches(2.95)
    text(s, MARGIN, y, BODY_W, Inches(0.34), "합칠 수 없는 이유", size=19, bold=True)
    code(s, MARGIN, y + Inches(0.44), BODY_W,
         ["지금            stage  [inner.read ── add 370us ──]  여럿이 동시에",
          "                publish            [inner.read + held.write]  ~0.2us 만 배타",
          "",
          "하나로 합치면   publish 가 쓰기 락 → 진행 중인 370us 삽입 전부를 기다린다",
          "                → 워커 N개의 처리량이 N분의 1"])

    # ---------------------------------------------------------------- 8
    s = page(prs, "검색은 락을 하나도 잡지 않는다", eyebrow="읽기",
             kicker="노드당 하는 일이 태그 비트 검사 하나다.")
    y = BODY_TOP + Inches(0.15)
    y = code(s, MARGIN, y, BODY_W,
             ["Some(matches) => self.matches(&index, query, k, |key| {",
              "    !held::is_staged(key) && matches(key)",
              "}),"]) + Inches(0.3)
    head = ["", "노드당", "통과 노드당", "질의당"]
    rows = [("예전", "읽기 락 + 집합 조회\n+ Arc<str> 할당", "map_elem_get\n+ refcount 보유", "held 가드"),
            ("지금", "비트 검사 하나", "ATTR ≤128B 복사", "스탬프 로드\n+ 슬롯 CAS")]
    colx = [MARGIN, MARGIN + Inches(1.3), MARGIN + Inches(5.0), MARGIN + Inches(8.4)]
    colw = [Inches(1.2), Inches(3.6), Inches(3.3), Inches(3.1)]
    for i, h in enumerate(head):
        text(s, colx[i], y, colw[i], Inches(0.3), h, size=11.5, font=MONO,
             color=MUTED, bold=True)
    rule(s, y + Inches(0.36))
    for j, row in enumerate(rows):
        yy = y + Inches(0.5) + j * Inches(0.78)
        for i, cell in enumerate(row):
            col = MOSS if (j == 1 and i == 0) else (FAIL if (j == 0 and i == 0) else INK)
            text(s, colx[i], yy, colw[i], Inches(0.6), cell.split("\n"), size=13.5,
                 bold=i == 0, color=col if i == 0 else INK, spacing=1.2)
        rule(s, yy + Inches(0.62), color=SOFT)

    y += Inches(2.35)
    text(s, MARGIN, y, BODY_W, Inches(1.0),
         [[("술어의 답이 결과 집합 전체를 좌우한다. ", {"bold": True}),
           ("잘못된 true 하나가 usearch의 가지치기 반경을 좁혀 진짜 이웃을 후보에서 자른다 — "
            "그래서 이 읽기는 반드시 유효해야 한다.", {})]],
         size=15, spacing=1.45)

    # ---------------------------------------------------------------- 9
    s = page(prs, "그러면 읽는 중에 해제되는 건 무엇이 막나", eyebrow="지연 회수")
    y = BODY_TOP - Inches(0.05)
    opts = [("읽는 동안 refcount로 고정", "불가", "engine.h에 acquire가 없다. 올리는 길이 map_elem_get뿐인데 필드가 필요해 순환이고 전역 캐시 락", FAIL),
            ("읽는 동안 락을 든다", "가능", "예전 방식. 노드당 락", MUTED),
            ("읽을 수 있는 자가 없음을 증명한다", "지금", "해제를 미룬다", MOSS)]
    for i, (a, verdict, why, col) in enumerate(opts):
        yy = y + i * Inches(0.86)
        text(s, MARGIN, yy, Inches(5.3), Inches(0.32), a, size=16, bold=True)
        rect(s, MARGIN + Inches(5.5), yy - Inches(0.02), Inches(0.95), Inches(0.34),
             fill=col, edge=None)
        label(s, MARGIN + Inches(5.5), yy - Inches(0.02), Inches(0.95), Inches(0.34),
              verdict, size=11.5, font=MONO, bold=True, color=SURFACE)
        text(s, MARGIN + Inches(6.7), yy + Inches(0.02), BODY_W - Inches(6.7), Inches(0.6),
             why, size=13, color=MUTED, spacing=1.3)

    y += Inches(2.75)
    y = code(s, MARGIN, y, BODY_W,
             ["검색 시작   스탬프를 슬롯에 꽂는다     슬롯 128개 고정 배열, 자기 홈에서 CAS",
              "방문 노드   그냥 역참조                락 없음",
              "검색 끝     슬롯을 비운다",
              "",
              "retire      (주소, 그때 스탬프 + 1) 을 큐에",
              "reclaim     아직 꽂힌 스탬프의 최솟값이 그 값 이상이면 반납"]) + Inches(0.24)
    text(s, MARGIN, y, BODY_W, Inches(0.4),
         [[("반납 조건 = ", {}), ("이 삭제보다 먼저 시작한 조회가 전부 끝났다", {"bold": True})]],
         size=16)

    # ---------------------------------------------------------------- 10
    s = page(prs, "끝난 개수로 세면 틀린다", eyebrow="지연 회수 · 반례",
             kicker="검색은 시작 순서대로 끝나지 않는다.")
    y = BODY_TOP + Inches(0.25)
    y = code(s, MARGIN, y, BODY_W,
         ["A 시작        started = 1",
          "B 시작        started = 2",
          "retire        스냅숏 = 2",
          "B 종료        ended = 1",
          "C 시작·종료   ended = 2   → 반납.  그런데 A가 아직 그 주소를 읽고 있다"]) + Inches(0.32)
    text(s, MARGIN, y, Inches(6.1), Inches(1.6),
         [[("그래서 개수가 아니라 아직 도는 것 중 가장 오래된 스탬프를 본다.", {"bold": True})],
          "슬롯이 다 차서 등록하지 못한 검색이 있으면 0을 돌려주어 회수를 전부 막는다 — "
          "드물고, 안전한 쪽이다."],
         size=15, spacing=1.45)
    code(s, MARGIN + Inches(6.7), y - Inches(0.15), BODY_W - Inches(6.7),
         ["fn oldest_reader(&self) -> u64 {",
          "    if self.unslotted.load(Acquire) > 0 {",
          "        return 0;",
          "    }",
          "    self.readers.iter()",
          "        .map(|c| c.load(Acquire))",
          "        .min()",
          "        .unwrap_or(NO_READER)",
          "}"])

    # ---------------------------------------------------------------- 11
    s = page(prs, "usearch는 우리 수명을 모른다", eyebrow="왜 우리가 지켜야 하나")
    y = BODY_TOP
    text(s, MARGIN, y, BODY_W, Inches(0.9),
         ["usearch는 자기 문서대로 스레드 안전하다 — 동시 add·search·remove가 자기 자료구조를 깨뜨리지 않는다. "
          "그런데 우리 키는 다른 시스템의 포인터다."],
         size=16, spacing=1.45)
    y += Inches(1.0)
    y = code(s, MARGIN, y, BODY_W,
             ["// index_dense.hpp — filtered_search 가 넘기는 술어",
              "auto allow = [free_key_copy, &predicate](member_cref_t const& member) noexcept {",
              "    return (vector_key_t)member.key != free_key_copy && predicate(member.key);",
              "};                                              // ↑ 락 없음, 평범한 필드 읽기"]) + Inches(0.3)
    text(s, MARGIN, y, Inches(6.1), Inches(2.0),
         [[("삭제가 콜백 도중에 들어옵니다.", {"bold": True})],
          "search_가 잡는 것은 스레드 컨텍스트 슬롯 하나뿐이고, remove는 slot_lookup_mutex_와 "
          "free_keys_mutex_를 잡는다. 겹치는 락이 없다.",
          "테스트로 못 박았다 — 콜백이 자기가 어느 노드 안인지 알리고 블록한 상태에서, "
          "다른 스레드의 remove가 완료된다."],
         size=14.5, spacing=1.4)
    rect(s, MARGIN + Inches(6.7), y - Inches(0.1), BODY_W - Inches(6.7), Inches(1.9),
         fill=SURFACE, edge=MOSS)
    text(s, MARGIN + Inches(6.95), y + Inches(0.12), BODY_W - Inches(7.2), Inches(1.5),
         [[("usearch가 지키는 것", {"bold": True, "color": MOSS})],
          "그래프 노드 · slot_lookup_ 표 · 동시 add/search/remove",
          [("우리가 지켜야 하는 것", {"bold": True, "color": FAIL})],
          "키가 가리키는 엔진 원소의 수명"],
         size=13.5, spacing=1.35)

    # ---------------------------------------------------------------- 12
    s = page(prs, "만료·축출된 인덱스는 누가 치우나", eyebrow="청소",
             kicker="일은 스레드가 하고, 커넥션은 쿠키가 필요한 호출 하나만 빌려준다.")
    y = BODY_TOP + Inches(0.2)
    y = code(s, MARGIN, y, BODY_W,
             ["arcvector-sweep   가장 차가운 10개 이름을 골라 내놓는다   ← 전체 엔트리를 읽는다",
              "  커넥션 ①        getattr(이름₁) → 없다, 돌려준다        ← 이거 하나. 끝",
              "  커넥션 ②        getattr(이름₂) → 있다, 잊는다",
              "arcvector-sweep   없다고 온 것을 빼고 해체한다"]) + Inches(0.3)
    text(s, MARGIN, y, Inches(6.1), Inches(1.9),
         [[("기준은 하나 — 쿠키가 필요한가.", {"bold": True})],
          "getattr은 키를 받는 API라 ACTION_BEFORE_READ가 쿠키를 넘긴다. "
          "마이그레이션 중이면 널 역참조다. 고르기와 해체는 쿠키가 필요 없다.",
          [("차가운 순서인 이유: 방금 만진 인덱스는 방금 Map이 있었던 인덱스다.", {})]],
         size=14.5, spacing=1.4)
    rect(s, MARGIN + Inches(6.7), y - Inches(0.12), BODY_W - Inches(6.7), Inches(1.75),
         fill=SURFACE, edge=FAIL)
    text(s, MARGIN + Inches(6.95), y + Inches(0.1), BODY_W - Inches(7.2), Inches(1.4),
         [[("해체를 레지스트리 락 안에서 하고 있었다", {"bold": True, "color": FAIL})],
          "마지막 Arc가 떨어지면 벡터마다 refcount 반납 + HNSW 소멸자. "
          "200만이면 그 시간이 전부 쓰기 락 안에서 흐르고, 모든 명령의 이름 조회가 선다."],
         size=13.5, spacing=1.35)

    # ---------------------------------------------------------------- 13
    s = page(prs, "데몬은 죽지 않는다", eyebrow="I5",
             kicker="메모리가 없으면 응답으로 답한다. abort하지 않는다.")
    y = BODY_TOP + Inches(0.2)
    cases = [
        ("자라는 컬렉션", "try_reserve 먼저, 실패는 에러로", MOSS),
        ("와이어에서 온 수", "VSIM VECTOR ix 4000000000 … 한 줄이 데몬을 죽였다. "
                            "파싱이 MAX_RESULTS로 막고, 버퍼는 실제 히트 수에서 잡는다", FAIL),
        ("reserve 결과를 버리는 것", "let _ = reserve() 뒤의 insert가 abort였다", FAIL),
        ("단일 할당", "format!·Box::new·collect는 여전히 abort. 미해결 §7", MUTED),
    ]
    for i, (a, b, col) in enumerate(cases):
        yy = y + i * Inches(1.0)
        rect(s, MARGIN, yy, Inches(0.09), Inches(0.78), fill=col, edge=None)
        text(s, MARGIN + Inches(0.35), yy, Inches(3.4), Inches(0.32), a, size=16, bold=True)
        text(s, MARGIN + Inches(3.95), yy, BODY_W - Inches(3.95), Inches(0.7),
             b, size=14, color=MUTED, spacing=1.35)

    # ---------------------------------------------------------------- 14
    s = page(prs, "남은 것", eyebrow="알려진 구멍",
             kicker="고칠 수 없다고 판단한 것들. 리뷰에서 다시 발견해도 새 버그가 아니다.")
    y = BODY_TOP + Inches(0.2)
    holes = [("mop update", "클라이언트가 같은 길이로 갱신하면 엔진이 제자리 memcpy를 한다. "
                            "그때 술어가 읽고 있으면 찢어진 값을 본다"),
             ("mop delete", "Map이 멀쩡해서 청소도 메타데이터 판정도 안 걸린다. "
                            "그 refcount는 인덱스가 드롭될 때까지 남는다"),
             ("HeldSet 중복", "usearch의 slot_lookup_과 같은 정보. 없애려면 export_keys를 "
                              "바인딩에 뚫어야 하고, 절약은 2M 기준 전체의 0.5%"),
             ("단일 할당", "해체 경로의 16MB Vec 둘이 가장 크다. 청크로 끊으면 없어진다"),
             ("빈 인덱스 1.1MB", "THREAD_SLOTS=64가 usearch에 스트라이프 락 테이블을 잡게 한다. "
                                 "인덱스 1000개면 1.1GB")]
    for i, (a, b) in enumerate(holes):
        yy = y + i * Inches(0.86)
        text(s, MARGIN, yy, Inches(2.9), Inches(0.32), a, size=15, font=MONO,
             bold=True, color=INK)
        text(s, MARGIN + Inches(3.1), yy, BODY_W - Inches(3.1), Inches(0.7),
             b, size=14, color=MUTED, spacing=1.35)
        rule(s, yy + Inches(0.66), color=SOFT)

    # ---------------------------------------------------------------- 15
    s = page(prs, "정리", eyebrow="한 문장씩")
    y = BODY_TOP + Inches(0.1)
    ends = [
        "Map이 원본, 그래프는 캐시. 실패는 캐시가 낡는 쪽으로 낸다.",
        "키를 주소로 두어 매핑 165MB를 없앴다. 대가로 refcount를 든다.",
        "쓰기는 두 저장소를 한 락 안에서 묶고, 실패할 수 있는 일은 전부 그 앞에 둔다.",
        "읽기는 락을 잡지 않는다. 해제를 미뤄서 같은 것을 보장한다.",
        "커넥션이 하는 일은 쿠키가 필요한 호출뿐. 나머지는 스레드로 보냈다.",
    ]
    for i, line in enumerate(ends):
        yy = y + i * Inches(0.82)
        text(s, MARGIN, yy, Inches(0.5), Inches(0.4), f"0{i + 1}", size=15,
             font=MONO, color=MOSS, bold=True)
        text(s, MARGIN + Inches(0.75), yy, BODY_W - Inches(0.75), Inches(0.6),
             line, size=17, spacing=1.35)

    text(s, MARGIN, H - Inches(1.25), BODY_W, Inches(0.5),
         "docs/내부구조.md · docs/코드리뷰.md · docs/미해결.md",
         size=13, font=MONO, color=MUTED)

    prs.save(out)
    print(f"wrote {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "ArcVector-내부설계.pptx")
